#!/usr/bin/env python3
"""Build a lecture deck (and live-coding files) from a lecture spec (YAML).

The spec usually lives in Lectures/src/week-NN/lecture.yaml and points to the course
settings in Lectures/config.yaml (key `config:`; found automatically if omitted).

Outputs:
  deck       meta.deck_file (relative to the spec), e.g. ../../5_Royxatlar.pptx
             default: <spec dir>/build/<slug>.pptx
  livecode/  <spec dir>/livecode/starter and solution (if the spec has `livecode`)
  build/lint.txt  design-rule warnings

An existing deck is never overwritten unless --force is given.

Usage (from the project root):
  .venv/Scripts/python.exe .claude/skills/dasturlash-ppt/scripts/build_lecture.py SPEC [--lint-only] [--force]
      [--deck PATH] [--only deck,code]
"""
import argparse
import itertools
import math
import re
import sys
from pathlib import Path

for _stream in (sys.stdout, sys.stderr):  # Windows consoles default to cp1252
    try:
        _stream.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass

try:
    import yaml
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    from lxml import etree
except ImportError as e:  # pragma: no cover
    sys.exit(f"Missing dependency ({e}). Install into the project venv:\n"
             "  .venv/Scripts/python.exe -m pip install python-pptx pygments pyyaml pillow")

sys.path.insert(0, str(Path(__file__).resolve().parent))
import companions  # noqa: E402  (live-coding files)

# --------------------------------------------------------------------------- constants

W, H = 13.333, 7.5          # 16:9 inches
M = 0.7                     # side margin
CW = W - 2 * M              # content width
KICKER_Y, HEAD_Y, HEAD_H = 0.45, 0.82, 1.35
CT, CB = 2.35, 6.65         # content top / bottom
FOOT_Y = 6.98

# Projector-friendly default: black on white for content, navy/green mode slides,
# one strong blue accent and an amber highlight that survive washed-out projectors.
DEFAULT_THEME = {
    "ink": "0B1F3A", "mode": "0F4C3A", "paper": "FFFFFF", "text": "111111",
    "muted": "4B5563", "accent": "0052CC", "signal": "FFC20E", "tint": "EEF2F7",
    "code_bg": "111827", "code_text": "E6E9EF", "on_dark": "FFFFFF", "on_dark_muted": "C3CCDA",
    "head_font": "Arial", "body_font": "Arial", "mono_font": "Courier New",
}

DEFAULT_MINUTES = {
    "title": 1, "retrieval": 4, "hook": 3, "objectives": 2, "section": 0.5, "assertion": 2.5,
    "code": 4, "predict": 4, "quiz": 5, "livecode": 10, "activity": 10, "compare": 3,
    "recap": 3, "exit": 3, "break": 4,
}
ACTIVE_TYPES = {"retrieval", "predict", "quiz", "livecode", "activity", "exit"}

# Every fixed text that appears on slides or in speaker notes.
STRINGS = {
    "uz": {
        "source": "Manba: {}",
        "min_chip": "{} DAQ",
        "notes_header": "⏱ {a}–{b}-daqiqa ({m} daq)",
        "notes_clicks": " · {} ta bosish",
        "notes_source": "Manba: {}",
        "break_title": "Qisqa tanaffus",
        "break_msg": "{} daqiqadan keyin davom etamiz",
        "objectives_kicker": "Bugun",
        "objectives_head": "Bugun nimani o'rganamiz",
        "retrieval_kicker": "Takrorlash",
        "retrieval_head": "O'tgan darsdan uchta savol",
        "hook_kicker": "Muammo",
        "predict_kicker": "Savol · natijani toping",
        "predict_head": "Bu kod nimani chiqaradi?",
        "predict_blank": "Javobingizni yozing, keyin yoningizdagi bilan solishtiring.",
        "answer": "Javob:",
        "quiz_kicker": "Savol · o'ylab ko'ring",
        "livecode_kicker": "Jonli kodlash · avval kuzating, keyin o'zingiz",
        "starter": "boshlang'ich fayl: livecode/starter/{}",
        "activity_kicker": "Navbat sizda",
        "activity_pairs": " · juftlikda",
        "open_file": "oching: livecode/starter/{}",
        "recap_kicker": "Xulosa",
        "recap_head": "Nimani eslab qolish kerak",
        "exit_kicker": "Chiqish savoli · 3 daqiqa",
        "exit_head": "Ketishdan oldin qog'ozga javob yozing",
        "p_retrieval": ("Talabalar daftarsiz, xotiradan javob yozadi (1–2 daqiqa). Keyin har bosishda bitta "
                        "javobni oching. Eslab chiqarish qayta o'qishdan ko'ra yaxshiroq yodda qoldiradi."),
        "p_hook": "Talabalarga 30–60 soniya o'ylab, taxmin aytish imkonini bering. Javobni hozir aytmang.",
        "p_code": "Kodni yuqoridan pastga tushuntiring; har bir ajratilgan joyda “agar ... bo'lsa nima bo'ladi?” deb so'rang.",
        "p_predict": ("Taxmin → juftlik → javob: 1 daqiqa har kim o'z taxminini yozadi, 1 daqiqa yonidagi "
                      "bilan solishtiradi. 2–3 talabadan javob so'rang, keyin bosib javobni oching va kodni "
                      "jonli ishga tushiring."),
        "p_misconception": "Tekshiriladigan noto'g'ri tasavvur: {}",
        "p_quiz": ("Peer instruction: 1) har kim alohida ovoz beradi (qo'l ko'tarish). 2) To'g'ri javob 30–70% "
                   "bo'lsa, 2 daqiqa yonidagi bilan muhokama (“sherigingizni ishontiring”), keyin qayta "
                   "ovoz. 70% dan ko'p bo'lsa, qisqa tushuntirib davom eting; 30% dan kam bo'lsa, qayta "
                   "tushuntiring. 3) Bosib javobni oching."),
        "p_distractors": "Noto'g'ri variantlar nega jalb qiladi: {}",
        "p_livecode": ("Jonli kodlash: boshlang'ich fayldan boshlang, har bir qarorni ovoz chiqarib tushuntiring, "
                       "tez-tez ishga tushiring, bitta haqiqiy xato qilib, uni ovoz chiqarib tuzating. "
                       "Maqsadlarga yetganda to'xtang va mashqqa o'ting."),
        "p_activity": "Qatorlar orasida yurib, ekranlarga qarang. Vaqtning yarmida qiynalayotganlarga maslahatni ko'rsating.",
        "p_hint": "Maslahat: {}",
        "p_recap": "Vaqt bo'lsa, har bir bandni ochishdan oldin talabalardan aytib berishni so'rang.",
        "p_exit": "Javoblarni yig'ib oling; keyingi darsdan oldin o'qing va darsni eng ko'p uchragan bo'shliqdan boshlang.",
    },
    "en": {
        "source": "Source: {}",
        "min_chip": "{} MIN",
        "notes_header": "⏱ min {a} → {b} ({m} min)",
        "notes_clicks": " · {} clicks",
        "notes_source": "Source: {}",
        "break_title": "Short break",
        "break_msg": "Back in {} minutes",
        "objectives_kicker": "Today",
        "objectives_head": "By the end of today you can…",
        "retrieval_kicker": "Warm-up · retrieval practice",
        "retrieval_head": "Three questions from last lecture",
        "hook_kicker": "The problem",
        "predict_kicker": "Question · predict the output",
        "predict_head": "What does this print?",
        "predict_blank": "Write your prediction down, then compare with a neighbour.",
        "answer": "Answer:",
        "quiz_kicker": "Question · concept check",
        "livecode_kicker": "Live coding · watch, then you try",
        "starter": "starter: livecode/starter/{}",
        "activity_kicker": "Your turn",
        "activity_pairs": " · in pairs",
        "open_file": "open: livecode/starter/{}",
        "recap_kicker": "Recap",
        "recap_head": "What to remember",
        "exit_kicker": "Exit ticket · 3 minutes",
        "exit_head": "Before you leave, answer on paper",
        "p_retrieval": "Students answer from memory first (no notes), 1–2 min, then reveal one answer per click.",
        "p_hook": "Let students think or guess for 30–60 s before you continue; don't answer it yet.",
        "p_code": "Narrate the code top-down; ask “what would happen if…?” at each highlight.",
        "p_predict": ("Predict → pair → reveal: 60 s to write a prediction, 60 s to compare with a "
                      "neighbour, take 2–3 answers, then click to reveal and run the code live."),
        "p_misconception": "Targets misconception: {}",
        "p_quiz": ("Peer instruction: vote alone; if 30–70% correct, 2 min neighbour discussion and re-vote; "
                   ">70%: explain briefly; <30%: re-teach. Then click to reveal."),
        "p_distractors": "Why the distractors are tempting: {}",
        "p_livecode": ("Live coding: start from the starter file, narrate decisions, run often, make one realistic "
                       "mistake and debug it aloud. Stop at the goals; hand over to the activity."),
        "p_activity": "Circulate and look at screens. At half time, show the hint to anyone stuck.",
        "p_hint": "Hint: {}",
        "p_recap": "Ask students to say each point back before you reveal it, if time allows.",
        "p_exit": "Collect answers; read them before next lecture and open with the most common gap.",
    },
}

# Syntax colours (dark code panel), keyed by pygments token type prefix
TOKEN_COLORS = [
    ("Comment", "9AA5B8", True), ("Keyword", "C792EA", False), ("Name.Builtin", "FFCB6B", False),
    ("Name.Function", "82AAFF", False), ("Name.Class", "FFCB6B", False), ("Name.Decorator", "82AAFF", False),
    ("Literal.String", "A5E3A1", False), ("Literal.Number", "F78C6C", False), ("Operator.Word", "C792EA", False),
    ("Operator", "89DDFF", False), ("Punctuation", "C3CCDA", False),
]

P_NS = "http://schemas.openxmlformats.org/presentationml/2006/main"

# Typographic apostrophes students and slides often mix; the course writes o' and g'.
APOSTROPHES = "‘’ʻʼ`"
SMART_QUOTES = "“”„"


# --------------------------------------------------------------------------- spec + config loading

def find_config(start: Path):
    for d in [start, *start.parents][:5]:
        for cand in (d / "config.yaml", d / "Lectures" / "config.yaml"):
            if cand.is_file():
                return cand
    return None


def config_to_spec(cfg: dict) -> dict:
    """Translate the teacher-facing config.yaml into spec defaults (meta/theme)."""
    lec, uni, crs = cfg.get("lecturer") or {}, cfg.get("university") or {}, cfg.get("course") or {}
    tch, ts, sl = cfg.get("teaching") or {}, cfg.get("title_slide") or {}, cfg.get("slides") or {}
    meta = {
        "lang": tch.get("language", "uz"),
        "language": tch.get("code_language", "python"),
        "duration_min": tch.get("lecture_minutes", 80),
        "course": crs.get("name", ""),
        "university": uni.get("name", ""),
        "university_short": uni.get("short", ""),
        "group": crs.get("group", ""),
        "academic_year": crs.get("academic_year", ""),
        "website": crs.get("website", ""),
        "lecturer": " ".join(x for x in (lec.get("title"), lec.get("name")) if x),
        "source": tch.get("default_source", ""),
        "apostrophe": tch.get("apostrophe", "'"),
        "transitions": sl.get("transitions", "fade"),
        "builds": sl.get("click_builds", True),
        "line_numbers": sl.get("line_numbers", True),
        "title_slide": {"university": ts.get("show_university", True), "group": ts.get("show_group", True),
                        "year": ts.get("show_year", True), "website": ts.get("show_website", False),
                        "email": ts.get("show_email", False), "email_value": lec.get("email", "")},
        "books": {b["id"]: b.get("cite", b["id"]) for b in (cfg.get("books") or []) if isinstance(b, dict) and b.get("id")},
    }
    theme = {}
    for k, v in (sl.get("colors") or {}).items():
        if v:
            theme[k] = str(v)
    fonts = sl.get("fonts") or {}
    for src, dst in (("heading", "head_font"), ("body", "body_font"), ("code", "mono_font")):
        if fonts.get(src):
            theme[dst] = fonts[src]
    return {"meta": meta, "theme": theme}


def merge(base: dict, over: dict) -> dict:
    out = dict(base)
    for k, v in over.items():
        if isinstance(v, dict) and isinstance(base.get(k), dict):
            out[k] = {**base[k], **v}
        else:
            out[k] = v
    return out


def load_spec(path: Path) -> dict:
    data = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    base = path.parent
    dirs = [base]
    if "extends" in data:                      # optional extra defaults file
        parent_path = (base / data.pop("extends")).resolve()
        parent = load_spec(parent_path)
        dirs += parent.get("_dirs", [])
        data = merge(parent, data)
    cfg_path = (base / data.pop("config")).resolve() if data.get("config") else find_config(base)
    if cfg_path and cfg_path.is_file():
        cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8")) or {}
        data = merge(config_to_spec(cfg), data)
        data["_config"] = str(cfg_path)
    data["_dirs"] = dirs
    data["_spec"] = path
    return data


def slugify(s: str) -> str:
    s = re.sub(r"[^\w\s-]", "", s.lower()).strip()
    return re.sub(r"[\s_-]+", "-", s)[:60] or "lecture"


def parse_lines(spec) -> list:
    """[2, '4-6', 8] or '2,4-6' -> [2,4,5,6,8]"""
    if spec is None:
        return []
    if isinstance(spec, (int, str)):
        spec = [spec]
    out = []
    for item in spec:
        for part in str(item).split(","):
            part = part.strip()
            if not part:
                continue
            if "-" in part:
                a, b = part.split("-", 1)
                out.extend(range(int(a), int(b) + 1))
            else:
                out.append(int(part))
    return out


# --------------------------------------------------------------------------- text helpers

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")
WIDTH = {"prop": 0.53}   # average glyph width in em for the body font (set per deck)


def runs_of(text: str):
    """Split inline markup into runs: [(text, style)] style in {None,'b','code'}."""
    out = []
    for part in INLINE.split(str(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            out.append((part[2:-2], "b"))
        elif part.startswith("`") and part.endswith("`"):
            out.append((part[1:-1], "code"))
        else:
            out.append((part, None))
    return out


def plain(text) -> str:
    return "".join(t for t, _ in runs_of(str(text)))


def est_lines(text: str, width_in: float, pt: float, mono=False) -> int:
    """Estimated wrapped line count; `code` spans count at monospace width and don't break."""
    prop = 0.60 if mono else WIDTH["prop"]
    lines = 0
    for para in str(text).split("\n"):
        words = []                                   # (width in inches) per unbreakable chunk
        for txt, style in runs_of(para):
            if style == "code" and len(txt) <= 26:
                words.append(len(txt) * 0.60 * pt / 72.0)
            else:
                f = 0.60 if style == "code" else prop
                words += [len(w) * f * pt / 72.0 for w in txt.split(" ") if w]
        space = prop * pt / 72.0
        n, cur = 1, 0.0
        for w in words:
            need = w if cur == 0 else cur + space + w
            if need <= width_in or cur == 0:
                cur = need
            else:
                n, cur = n + 1, w
        lines += n
    return lines


def fit_size(paras, width_in, height_in, start, floor, line=1.2, gap_pt=0.0, mono=False):
    """Largest size (pt) from start down to floor at which paras fit. Returns (size, fits)."""
    size = start
    while size >= floor:
        h = 0.0
        for p in paras:
            n = est_lines(p, width_in, size, mono)
            h += n * size * line / 72.0 + gap_pt / 72.0
        if h <= height_in:
            return size, True
        size -= 1
    return floor, False


# --------------------------------------------------------------------------- animation XML

class Timing:
    """Collects click groups for one slide and writes <p:timing>."""

    def __init__(self):
        self.groups = []          # list of list of (spid, mode, para, is_sp)

    def click(self, *entries):
        group = []
        for e in entries:
            if e is None:
                continue
            shape, mode, para = (e + (None,))[:3] if isinstance(e, tuple) else (e, "in", None)
            if shape is None:
                continue
            is_sp = shape._element.tag == qn("p:sp")
            group.append((shape.shape_id, mode, para, is_sp))
        if group:
            self.groups.append(group)

    def __len__(self):
        return len(self.groups)

    def xml(self, fade=True):
        if not self.groups:
            return None
        ids = itertools.count(3)
        grp_of, next_grp, bld, clicks = {}, {}, [], []
        for group in self.groups:
            c_outer, c_inner = next(ids), next(ids)
            effects = []
            for j, (spid, mode, para, is_sp) in enumerate(group):
                key = (spid, mode, "p" if para is not None else "w")
                if key not in grp_of:
                    grp_of[key] = next_grp.get(spid, 0)
                    next_grp[spid] = grp_of[key] + 1
                    if is_sp:
                        bld.append((spid, grp_of[key], "p" if para is not None else "whole"))
                grp = grp_of[key]
                node = "clickEffect" if j == 0 else "withEffect"
                txel = f'<p:txEl><p:pRg st="{para}" end="{para}"/></p:txEl>' if para is not None else ""
                tgt = f'<p:tgtEl><p:spTgt spid="{spid}">{txel}</p:spTgt></p:tgtEl>'
                vis = '<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>'
                if mode == "in":
                    c1, c2 = next(ids), next(ids)
                    s = (f'<p:par><p:cTn id="{c1}" presetID="{10 if fade else 1}" presetClass="entr" presetSubtype="0" '
                         f'fill="hold" grpId="{grp}" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst>'
                         f'<p:childTnLst><p:set><p:cBhvr><p:cTn id="{c2}" dur="1" fill="hold"><p:stCondLst>'
                         f'<p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}{vis}</p:cBhvr><p:to><p:strVal val="visible"/>'
                         f'</p:to></p:set>')
                    if fade:
                        c3 = next(ids)
                        s += (f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{c3}" dur="400"/>'
                              f'{tgt}</p:cBhvr></p:animEffect>')
                    s += '</p:childTnLst></p:cTn></p:par>'
                else:
                    c1, c2, c3 = next(ids), next(ids), next(ids)
                    s = (f'<p:par><p:cTn id="{c1}" presetID="10" presetClass="exit" presetSubtype="0" fill="hold" '
                         f'grpId="{grp}" nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>'
                         f'<p:animEffect transition="out" filter="fade"><p:cBhvr><p:cTn id="{c2}" dur="300"/>{tgt}'
                         f'</p:cBhvr></p:animEffect><p:set><p:cBhvr><p:cTn id="{c3}" dur="1" fill="hold"><p:stCondLst>'
                         f'<p:cond delay="299"/></p:stCondLst></p:cTn>{tgt}{vis}</p:cBhvr><p:to>'
                         f'<p:strVal val="hidden"/></p:to></p:set></p:childTnLst></p:cTn></p:par>')
                effects.append(s)
            clicks.append(f'<p:par><p:cTn id="{c_outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/>'
                          f'</p:stCondLst><p:childTnLst><p:par><p:cTn id="{c_inner}" fill="hold"><p:stCondLst>'
                          f'<p:cond delay="0"/></p:stCondLst><p:childTnLst>{"".join(effects)}</p:childTnLst></p:cTn>'
                          f'</p:par></p:childTnLst></p:cTn></p:par>')
        bld_xml = "".join(
            f'<p:bldP spid="{s}" grpId="{g}" build="p"/>' if b == "p" else f'<p:bldP spid="{s}" grpId="{g}" animBg="1"/>'
            for s, g, b in bld)
        return (f'<p:timing xmlns:p="{P_NS}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" '
                f'nodeType="tmRoot"><p:childTnLst><p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" '
                f'nodeType="mainSeq"><p:childTnLst>{"".join(clicks)}</p:childTnLst></p:cTn><p:prevCondLst>'
                f'<p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
                f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond>'
                f'</p:nextCondLst></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
                + (f'<p:bldLst>{bld_xml}</p:bldLst>' if bld_xml else "") + '</p:timing>')


def finish_slide_xml(slide, timing: Timing, transition: str):
    sld = slide._element
    for tag in ("p:transition", "p:timing"):
        old = sld.find(qn(tag))
        if old is not None:
            sld.remove(old)
    anchor = sld.find(qn("p:clrMapOvr"))
    idx = list(sld).index(anchor) + 1 if anchor is not None else len(sld)
    if transition == "fade":
        sld.insert(idx, etree.fromstring(f'<p:transition xmlns:p="{P_NS}" spd="med"><p:fade/></p:transition>'))
        idx += 1
    tx = timing.xml()
    if tx:
        sld.insert(idx, etree.fromstring(tx))


# --------------------------------------------------------------------------- deck builder

class LectureDeck:
    def __init__(self, spec: dict):
        self.spec = spec
        self.meta = spec.get("meta", {})
        self.t = {**DEFAULT_THEME, **{k: str(v).lstrip("#") for k, v in (spec.get("theme") or {}).items()}}
        WIDTH["prop"] = 0.50 if "calibri" in self.t["body_font"].lower() else 0.55
        self.s = STRINGS.get(str(self.meta.get("lang", "uz")), STRINGS["uz"])
        self.dirs = spec.get("_dirs", [Path(".")])
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = Inches(W), Inches(H)
        self.blank = self.prs.slide_layouts[6]
        self.warnings = []
        self.section = ""
        self.section_no = 0
        self.clock = 0.0
        self.since_active = 0.0
        self.n = 0
        self.timeline = []  # (n, type, start, end, clicks, label)
        self.lang = self.meta.get("language", "python")
        self.builds_default = bool(self.meta.get("builds", True))
        self.apostrophe = str(self.meta.get("apostrophe", "'"))

    # ---- primitives --------------------------------------------------------
    def warn(self, msg):
        self.warnings.append(f"slide {self.n}: {msg}")

    def fix_text(self, s):
        """Course style for visible text: o' and g' with one apostrophe character."""
        s = str(s)
        if self.apostrophe == "'":
            s = re.sub(rf"([oOgG])[{APOSTROPHES}]", r"\1'", s)
            s = re.sub(r"(\w)[’ʼ](\w)", r"\1'\2", s)   # ma’lumot -> ma'lumot
        return s

    def cite(self, source):
        """'@thinkpython, 3-bob' -> 'A. Downey. Think Python, 3-bob' using config books."""
        if not source:
            return source
        books = self.meta.get("books") or {}
        return re.sub(r"@([\w-]+)", lambda m: books.get(m.group(1), m.group(0)), str(source))

    def rgb(self, key_or_hex):
        h = self.t.get(key_or_hex, key_or_hex)
        return RGBColor.from_string(str(h).lstrip("#").upper())

    def bg(self, slide, key):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self.rgb(key)

    def rect(self, slide, x, y, w, h, fill=None, line=None, line_w=1.5, rounded=False, radius=0.06,
             alpha=None, name=None, shape=None):
        kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE)
        sp = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        style = sp._element.find(qn("p:style"))
        if style is not None:
            sp._element.remove(style)
        if rounded and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
            sp.adjustments[0] = radius
        if fill:
            sp.fill.solid()
            sp.fill.fore_color.rgb = self.rgb(fill)
            if alpha is not None:
                clr = sp._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                etree.SubElement(clr, qn("a:alpha")).set("val", str(int(alpha * 1000)))
        else:
            sp.fill.background()
        if line:
            sp.line.color.rgb = self.rgb(line)
            sp.line.width = Pt(line_w)
        else:
            sp.line.fill.background()
        if name:
            sp.name = name
        return sp

    def text(self, slide_or_shape, x, y, w, h, paras, size, color="text", font=None, bold=False, italic=False,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line=1.15, gap=0, margin=(0.05, 0.03),
             bullets=False, numbered=False, name=None, code_color=None, marker="accent"):
        """paras: list of str (inline markup); returns shape."""
        if hasattr(slide_or_shape, "shapes"):
            shp = slide_or_shape.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        else:
            shp = slide_or_shape
        tf = shp.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(margin[0])
        tf.margin_top = tf.margin_bottom = Inches(margin[1])
        font = font or self.t["body_font"]
        if isinstance(paras, str):
            paras = [paras]
        for i, ptxt in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.line_spacing = line
            p.space_after = Pt(gap)
            p.space_before = Pt(0)
            if bullets or numbered:
                self._bullet(p, numbered, marker)
            for txt, style in runs_of(ptxt):
                r = p.add_run()
                if style == "code":   # keep short code like `km = 10` on one line
                    r.text = txt.replace(" ", "\u00a0") if len(txt) <= 26 else txt
                else:
                    r.text = self.fix_text(txt)
                f = r.font
                f.size = Pt(size)
                f.bold = bold or style == "b"
                f.italic = italic
                if style == "code":
                    f.name = self.t["mono_font"]
                    f.color.rgb = self.rgb(code_color or "accent")
                else:
                    f.name = font
                    f.color.rgb = self.rgb(color)
        if name:
            shp.name = name
        return shp

    def _bullet(self, p, numbered, marker="accent"):
        pPr = p._p.get_or_add_pPr()
        ind = int(Inches(0.42 if not numbered else 0.5))
        pPr.set("marL", str(ind))
        pPr.set("indent", str(-ind))
        for tag in ("a:buClr", "a:buSzPct", "a:buFont", "a:buChar", "a:buAutoNum", "a:buNone"):
            for el in pPr.findall(qn(tag)):
                pPr.remove(el)
        clr = etree.SubElement(pPr, qn("a:buClr"))
        etree.SubElement(clr, qn("a:srgbClr")).set("val", self.t.get(marker, marker))
        if numbered:
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", self.t["body_font"])
            etree.SubElement(pPr, qn("a:buAutoNum")).set("type", "arabicPeriod")
        else:
            etree.SubElement(pPr, qn("a:buFont")).set("typeface", "Arial")
            etree.SubElement(pPr, qn("a:buChar")).set("char", "•")

    def new_slide(self, bg="paper"):
        s = self.prs.slides.add_slide(self.blank)
        self.bg(s, bg)
        return s

    def kicker(self, slide, label, color="accent"):
        self.text(slide, M, KICKER_Y, CW, 0.36, [self.fix_text(label).upper()], 14, color, font=self.t["mono_font"],
                  bold=True, margin=(0, 0), name="Kicker")

    def headline(self, slide, txt, color="text", y=HEAD_Y, h=HEAD_H, w=CW):
        txt = str(txt).rstrip().rstrip(".")
        size, fits = fit_size([txt], w - 0.1, h, 32, 26, line=1.05)
        if not fits:
            self.warn(f"headline too long for 2 lines; shorten it: “{plain(txt)[:60]}…”")
        words = len(plain(txt).split())
        if words > 16:
            self.warn(f"headline has {words} words (aim for ≤ 14)")
        return self.text(slide, M, y, w, h, [txt], size, color, font=self.t["head_font"], bold=True,
                         line=1.05, margin=(0, 0), name="Headline")

    def footer(self, slide, source=None, dark=False):
        col = "on_dark_muted" if dark else "muted"
        if source:
            self.text(slide, M, FOOT_Y, CW - 1.2, 0.32, [self.s["source"].format(self.cite(source))], 12, col,
                      margin=(0, 0), name="Source")
        self.text(slide, W - M - 1.0, FOOT_Y, 1.0, 0.32, [str(self.n)], 12, col, align=PP_ALIGN.RIGHT,
                  margin=(0, 0), name="Slide number")

    def chip(self, slide, x, y, label, fill="signal", color="text"):
        w = 0.36 + 0.14 * len(label)
        c = self.rect(slide, x, y, w, 0.44, fill=fill, rounded=True, radius=0.5, name="Chip")
        self.text(c, 0, 0, 0, 0, [label], 15, color, font=self.t["mono_font"], bold=True,
                  align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=(0.05, 0))
        return c

    def resolve(self, rel):
        for d in self.dirs:
            p = Path(d) / rel
            if p.exists():
                return p
        return None

    # ---- style checks (course rules) -----------------------------------------
    def check_text(self, *texts):
        for t in texts:
            if not t:
                continue
            low = plain(t).lower()
            if re.search(r"\brost\b|\byolg'on\b|\byolg’on\b", low):
                self.warn("write `True`/`False`, not rost/yolg'on (course rule)")

    def check_code(self, code, language):
        code = str(code)
        lang = (language or self.lang).lower()
        if any(q in code for q in SMART_QUOTES):
            self.warn("code contains typographic quotes “ ”; use straight quotes \"")
        if re.search("[‘’ʻʼ]", code):
            self.warn("code contains a typographic apostrophe (o‘, ’); use o' and g' in strings")
        if "\t" in code:
            self.warn("code contains tabs; use 4 spaces")
        if lang == "python":
            if ".title()" in code:
                self.warn("`.title()` breaks o'/g' (O'Tkir); use string.capwords() or .capitalize()")
            for m in re.finditer(r"def \w+\(([^)]*)\)", code):
                if re.search(r"\w\s+=\s*|\w\s*=\s+", m.group(1)):
                    self.warn(f"PEP 8: no spaces around = in default arguments: `{m.group(0)}`")
            for ln in code.split("\n"):
                if len(ln) > 79:
                    self.warn("PEP 8: code line longer than 79 characters")
                    break

    # ---- evidence renderers --------------------------------------------------
    def evidence(self, slide, ev, x, y, w, h, tm: Timing, build=True):
        kind = (ev or {}).get("kind", "bullets")
        fn = getattr(self, f"ev_{kind}", None)
        if fn is None:
            self.warn(f"unknown evidence kind '{kind}'")
            return
        fn(slide, ev, x, y, w, h, tm, build)

    def ev_bullets(self, slide, ev, x, y, w, h, tm, build):
        self.bullets(slide, ev.get("points", []), x, y, w, h, tm, build)

    def bullets(self, slide, points, x, y, w, h, tm, build, size=26):
        points = [str(p) for p in points or []]
        self.check_text(*points)
        if len(points) > 4:
            self.warn(f"{len(points)} bullets (aim for ≤ 4; split the slide)")
        for p in points:
            if len(plain(p).split()) > 12:
                self.warn(f"long bullet ({len(plain(p).split())} words): “{plain(p)[:50]}…”")
        sz, fits = fit_size(points, w - 0.55, h, size, 22, line=1.15, gap_pt=14)
        if not fits:
            self.warn("bullets overflow their area even at 22 pt; cut words")
        shp = self.text(slide, x, y, w, h, points, sz, "text", bullets=True, line=1.15, gap=14, name="Points")
        if build:
            for i in range(len(points)):
                tm.click((shp, "in", i))
        return shp

    def code_panel(self, slide, code, x, y, w, h, tm, steps=None, language=None, notes_box=None):
        """Dark code panel with syntax colouring; steps highlight lines.
        notes_box: (x,y,w,h) for accumulating step notes, or None."""
        code = str(code).rstrip("\n").replace("\t", "    ")
        self.check_code(code, language)
        lines = code.split("\n")
        show_ln = self.meta.get("line_numbers", True)
        gutter = len(str(len(lines))) + 2 if show_ln else 0
        maxlen = max((len(l) for l in lines), default=1) + gutter
        pad_x, pad_y = 0.28, 0.22
        avail_w, avail_h = w - 2 * pad_x, h - 2 * pad_y
        lh_factor = 1.22
        size = min(22.0, avail_h * 72 / (len(lines) * lh_factor), avail_w * 72 / (maxlen * 0.60))
        size = math.floor(size * 2) / 2
        if size < 16:
            self.warn(f"code shrinks to {size:.1f} pt ({len(lines)} lines, longest {maxlen - gutter} chars); "
                      "show fewer/shorter lines")
        if len(lines) > 14:
            self.warn(f"{len(lines)} lines of code (aim for ≤ 14)")
        size = max(size, 12)
        lh = size * lh_factor / 72.0
        used_h = len(lines) * lh + 2 * pad_y
        panel_h = min(h, max(used_h, 1.0))
        self.rect(slide, x, y, w, panel_h, fill="code_bg", rounded=True, radius=0.03, name="Code panel")

        step_shapes = []
        for st in steps or []:
            bars = []
            for ln in parse_lines(st.get("lines")):
                if 1 <= ln <= len(lines):
                    by = y + pad_y + (ln - 1) * lh
                    bars.append(self.rect(slide, x + 0.08, by, w - 0.16, lh, fill="signal", alpha=34,
                                          name=f"Highlight line {ln}"))
                else:
                    self.warn(f"step highlights line {ln} but code has {len(lines)} lines")
            step_shapes.append(bars)

        # line numbers in their own box, so the code text copies and extracts cleanly
        gut_w = gutter * 0.6 * size / 72.0 if show_ln else 0.0
        if show_ln:
            gb = slide.shapes.add_textbox(Inches(x + pad_x), Inches(y + pad_y), Inches(gut_w),
                                          Inches(panel_h - 2 * pad_y))
            gb.name = "Line numbers"
            gtf = gb.text_frame
            gtf.word_wrap = False
            gtf.auto_size = MSO_AUTO_SIZE.NONE
            for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
                setattr(gtf, side, 0)
            for i in range(len(lines)):
                p = gtf.paragraphs[0] if i == 0 else gtf.add_paragraph()
                p.line_spacing = Pt(size * lh_factor)
                p.space_after = p.space_before = Pt(0)
                r = p.add_run()
                r.text = f"{i + 1:>{gutter - 2}}"
                self._mono(r, size, "8792A8")
        tb = slide.shapes.add_textbox(Inches(x + pad_x + gut_w), Inches(y + pad_y), Inches(avail_w - gut_w),
                                      Inches(panel_h - 2 * pad_y))
        tb.name = "Code"
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        for i, line_tokens in enumerate(self._tokenize(code, language or self.lang)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = Pt(size * lh_factor)
            p.space_after = Pt(0)
            p.space_before = Pt(0)
            if not line_tokens:
                r = p.add_run()
                r.text = " "
                self._mono(r, size, self.t["code_text"])
            for txt, colour, ital in line_tokens:
                r = p.add_run()
                r.text = txt
                self._mono(r, size, colour, ital)

        note_shapes = []
        if steps and notes_box:
            nx, ny, nw, nh = notes_box
            notes = [str(s.get("note", "")) for s in steps]
            self.check_text(*notes)
            nsz, fits = fit_size(notes, nw - 0.7, nh, 22, 18, line=1.1, gap_pt=10)
            if not fits:
                self.warn("step notes overflow; shorten them")
            cy = ny
            for k, note in enumerate(notes):
                hh = est_lines(note, nw - 0.62, nsz) * nsz * 1.1 / 72 + 0.12
                badge = self.rect(slide, nx, cy + 0.02, 0.44, 0.44, fill="signal", shape=MSO_SHAPE.OVAL,
                                  name=f"Step {k + 1} badge")
                self.text(badge, 0, 0, 0, 0, [str(k + 1)], 16, "text", font=self.t["mono_font"], bold=True,
                          align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0))
                nt = self.text(slide, nx + 0.6, cy - 0.02, nw - 0.62, hh + 0.1, [note], nsz, "text",
                               line=1.1, margin=(0, 0), name=f"Step {k + 1} note")
                note_shapes.append((badge, nt))
                cy += max(hh, 0.5) + nsz * 0.6 / 72 + 0.12

        prev = []
        for k, bars in enumerate(step_shapes):
            entries = [(b, "in") for b in bars]
            if note_shapes:
                entries += [(note_shapes[k][0], "in"), (note_shapes[k][1], "in")]
            entries += [(b, "out") for b in prev]
            tm.click(*entries)
            prev = bars
        return panel_h

    def _mono(self, run, size, colour, italic=False):
        f = run.font
        f.name = self.t["mono_font"]
        f.size = Pt(size)
        f.italic = italic
        f.color.rgb = RGBColor.from_string(str(colour).lstrip("#").upper())

    def _tokenize(self, code, language):
        default = self.t["code_text"]
        try:
            from pygments.lexers import get_lexer_by_name
            toks = list(get_lexer_by_name(language, stripnl=False, ensurenl=False).get_tokens(code))
        except Exception:
            toks = [(None, code)]
        lines = [[]]
        for ttype, value in toks:
            colour, ital = default, False
            if ttype is not None:
                name = str(ttype).replace("Token.", "")
                for prefix, c, it in TOKEN_COLORS:
                    if name.startswith(prefix):
                        colour, ital = c, it
                        break
            for j, part in enumerate(value.split("\n")):
                if j > 0:
                    lines.append([])
                if part:
                    lines[-1].append((part, colour, ital))
        n = len(code.split("\n"))
        return (lines + [[]] * n)[:n]

    def ev_code(self, slide, ev, x, y, w, h, tm, build):
        self.code_panel(slide, ev.get("code", ""), x, y, w, h, tm, ev.get("steps"), ev.get("language"))

    def ev_output(self, slide, ev, x, y, w, h, tm, build):
        """Program output in a light terminal box (e.g. after a code slide)."""
        txt = str(ev.get("text", "")).rstrip("\n")
        lines = txt.split("\n")
        size = min(24.0, (h - 0.5) * 72 / (len(lines) * 1.25), (w - 0.6) * 72 / (max(map(len, lines)) * 0.6 + 1))
        size = max(14, math.floor(size))
        box_h = min(h, len(lines) * size * 1.25 / 72 + 0.5)
        box = self.rect(slide, x, y, w, box_h, fill="tint", line="muted", line_w=1, rounded=True, radius=0.04,
                        name="Output")
        tf = box.text_frame
        tf.word_wrap = False
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = tf.margin_right = Inches(0.3)
        tf.margin_top = tf.margin_bottom = Inches(0.22)
        for i, ln in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = ln or " "
            self._mono(r, size, self.t["text"])
        if build and ev.get("build", True):
            tm.click((box, "in"))

    def ev_image(self, slide, ev, x, y, w, h, tm, build):
        p = self.resolve(ev.get("path", ""))
        cap = ev.get("caption")
        cap_h = 0.4 if cap else 0
        if not p:
            self.warn(f"image not found: {ev.get('path')}")
            box = self.rect(slide, x, y, w, h - cap_h, fill="tint", rounded=True, name="Missing image")
            self.text(box, 0, 0, 0, 0, [f"[missing image: {ev.get('path')}]"], 18, "muted",
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
            return
        from PIL import Image
        with Image.open(p) as im:
            iw, ih = im.size
        if iw < 1000:
            self.warn(f"image {p.name} is only {iw}px wide; may look blurry when projected")
        bw, bh = w, h - cap_h
        scale = min(bw / iw, bh / ih)
        pw, ph = iw * scale, ih * scale
        px, py = x + (bw - pw) / 2, y + (bh - ph) / 2
        pic = slide.shapes.add_picture(str(p), Inches(px), Inches(py), Inches(pw), Inches(ph))
        alt = ev.get("alt")
        if not alt:
            self.warn(f"image {p.name} has no alt text")
            alt = cap or p.stem
        pic._element.nvPicPr.cNvPr.set("descr", plain(alt))
        pic.name = f"Figure {p.stem}"
        if cap:
            self.text(slide, x, py + ph + 0.06, w, 0.36, [cap], 14, "muted", italic=True, align=PP_ALIGN.CENTER,
                      margin=(0, 0), name="Caption")

    def ev_array(self, slide, ev, x, y, w, h, tm, build):
        cells = [str(c) for c in ev.get("cells", [])]
        n = max(len(cells), 1)
        if n > 12:
            self.warn(f"array has {n} cells (≤ 12 stays readable)")
        label = ev.get("label")
        steps = ev.get("steps") or []
        notes_h = 0.55 * len(steps) + (0.2 if steps else 0)
        cell = min(1.15, (w - 0.2) / n, max(0.6, (h - notes_h - 0.9)))
        total = cell * n
        x0 = x + (w - total) / 2
        top = y + (0.5 if label else 0.1)
        if label:
            self.text(slide, x0, y, total, 0.4, [label], 16, "muted", font=self.t["mono_font"], margin=(0, 0))
        longest = max((len(c) for c in cells), default=1)
        fsize = max(12, min(24, cell * 72 / (max(longest, 2) * 0.62)))
        for i, val in enumerate(cells):
            c = self.rect(slide, x0 + i * cell, top, cell, cell, fill="tint", line="muted", line_w=1.25,
                          name=f"Cell {i}")
            self.text(c, 0, 0, 0, 0, [val], fsize, "text", font=self.t["mono_font"], bold=True,
                      align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, margin=(0.02, 0))
            if ev.get("indices", True):
                self.text(slide, x0 + i * cell, top + cell + 0.05, cell, 0.35, [str(i)], 14, "muted",
                          font=self.t["mono_font"], align=PP_ALIGN.CENTER, margin=(0, 0))
        prev = []
        ny = top + cell + 0.55
        for k, st in enumerate(steps):
            overlays = []
            for ci in parse_lines(st.get("cells")):
                if 0 <= ci < n:
                    overlays.append(self.rect(slide, x0 + ci * cell, top, cell, cell, fill="signal", alpha=45,
                                              line="signal", line_w=3, name=f"Highlight cell {ci}"))
            note = None
            if st.get("note"):
                note = self.text(slide, x0, ny + k * 0.55, max(total, 6), 0.5, [f"{k + 1}.  {st['note']}"], 20,
                                 "text", margin=(0, 0), name=f"Step {k + 1} note")
            tm.click(*([(o, "in") for o in overlays] + ([(note, "in")] if note else []) + [(o, "out") for o in prev]))
            prev = overlays

    def ev_flow(self, slide, ev, x, y, w, h, tm, build):
        nodes = [str(n) for n in ev.get("nodes", [])]
        labels = [str(l) for l in ev.get("labels", [])]
        n = len(nodes)
        if n == 0:
            return
        if n > 5:
            self.warn(f"flow has {n} boxes (≤ 5 fits)")
        arrow_w = 0.55
        box_w = (w - arrow_w * (n - 1)) / n
        box_h = min(1.4, h * 0.55)
        by = y + (h - box_h) / 2 - (0.25 if labels else 0)
        longest = max(len(plain(t)) for t in nodes)
        fsize = max(16, min(26, box_w * 72 / (max(longest, 4) * 0.6) * 1.7))
        prev_arrow = None
        for i, label in enumerate(nodes):
            bx = x + i * (box_w + arrow_w)
            last = i == n - 1
            b = self.rect(slide, bx, by, box_w, box_h, fill="accent" if last else "tint", rounded=True, radius=0.12,
                          name=f"Flow box {i + 1}")
            self.text(b, 0, 0, 0, 0, [label], fsize, "on_dark" if last else "text", bold=True, align=PP_ALIGN.CENTER,
                      anchor=MSO_ANCHOR.MIDDLE, margin=(0.08, 0.04), code_color="on_dark" if last else "accent")
            cap = None
            if i < len(labels) and labels[i]:
                cap = self.text(slide, bx, by + box_h + 0.12, box_w, 0.5, [labels[i]], 18, "muted",
                                align=PP_ALIGN.CENTER, margin=(0, 0), name=f"Flow label {i + 1}")
            if build:
                tm.click(*([(prev_arrow, "in")] if prev_arrow else []), (b, "in"), (cap, "in") if cap else None)
            if i < n - 1:
                prev_arrow = self.rect(slide, bx + box_w + 0.1, by + box_h / 2 - 0.18, arrow_w - 0.2, 0.36,
                                       fill="muted", shape=MSO_SHAPE.RIGHT_ARROW, name=f"Arrow {i + 1}")

    def ev_table(self, slide, ev, x, y, w, h, tm, build):
        rows = [[str(c) for c in r] for r in ev.get("rows", [])]
        if not rows:
            return
        nr, nc = len(rows), max(len(r) for r in rows)
        row_h = min(0.7, h / nr)
        gf = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        gf.name = "Table"
        tbl = gf.table
        size = 22 if nr <= 6 else 18
        for r in range(nr):
            tbl.rows[r].height = Inches(row_h)
            for c in range(nc):
                cell = tbl.cell(r, c)
                val = rows[r][c] if c < len(rows[r]) else ""
                cell.fill.solid()
                cell.fill.fore_color.rgb = self.rgb("accent" if r == 0 else ("tint" if r % 2 == 0 else "paper"))
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                cell.text_frame.paragraphs[0].text = ""
                self.text(cell, 0, 0, 0, 0, [val], size, "on_dark" if r == 0 else "text", bold=(r == 0),
                          margin=(0.12, 0.03))

    def ev_terms(self, slide, ev, x, y, w, h, tm, build):
        """Term -> meaning rows (like `def` — yangi funksiya yaratadi), revealed one per click."""
        items = ev.get("items") or []
        n = max(len(items), 1)
        row = min(0.95, h / n)
        tw = min(4.2, w * 0.38)
        for i, it in enumerate(items):
            term, meaning = str(it.get("term", "")), str(it.get("means", ""))
            ty = y + i * row
            a = self.text(slide, x, ty, tw, row, [f"`{term}`"], 24, "text", margin=(0, 0), code_color="accent",
                          anchor=MSO_ANCHOR.TOP, name=f"Term {i + 1}")
            ms, _ = fit_size([meaning], w - tw - 0.3, row, 24, 18)
            b = self.text(slide, x + tw + 0.3, ty + 0.02, w - tw - 0.3, row, [meaning], ms, "text", margin=(0, 0),
                          name=f"Meaning {i + 1}")
            self.check_text(meaning)
            if build:
                tm.click((a, "in"), (b, "in"))

    # ---- slide types ---------------------------------------------------------
    def add(self, sd: dict):
        typ = sd.get("type", "assertion")
        fn = getattr(self, f"s_{typ}", None)
        if fn is None:
            self.warnings.append(f"after slide {self.n}: unknown slide type '{typ}' (skipped)")
            return
        self.n += 1
        minutes = float(sd.get("minutes", DEFAULT_MINUTES.get(typ, 2)))
        tm = Timing()
        self.check_text(sd.get("headline"), sd.get("question"), sd.get("title"))
        slide, protocol = fn(sd, tm)
        start, end = self.clock, self.clock + minutes
        self.clock = end
        if typ in ACTIVE_TYPES:
            self.since_active = 0
        else:
            self.since_active += minutes
            if self.since_active > 20 and typ not in ("title", "section", "break"):
                self.warn(f"{self.since_active:.0f} min without an active slide (predict/quiz/activity…)")
                self.since_active = -1e9
        fade = str(self.meta.get("transitions", "fade")) == "fade"
        finish_slide_xml(slide, tm, "fade" if fade else "none")
        source = sd.get("source") or (self.meta.get("source") if typ not in ("title", "section", "break") else None)
        header = self.s["notes_header"].format(a=f"{start:g}", b=f"{end:g}", m=f"{minutes:g}")
        if len(tm):
            header += self.s["notes_clicks"].format(len(tm))
        body = [header, ""]
        if sd.get("notes"):
            body.append(self.fix_text(str(sd["notes"]).strip()))
            self.check_text(sd["notes"])
        elif typ in ("assertion", "code", "compare", "hook"):
            self.warn("no speaker notes (what to explain, an example, a question to ask)")
        if protocol:
            body += ["", protocol]
        if source:
            body += ["", self.s["notes_source"].format(self.cite(source))]
        slide.notes_slide.notes_text_frame.text = "\n".join(body)
        self.timeline.append((self.n, typ, start, end, len(tm),
                              plain(sd.get("headline") or sd.get("title") or sd.get("question") or "")))

    def _content_head(self, s, sd, default_kicker=None, kicker_color="accent"):
        self.kicker(s, sd.get("kicker") or default_kicker or self.section or self.meta.get("title", ""), kicker_color)
        if sd.get("headline"):
            self.headline(s, sd["headline"])

    def _build(self, sd):
        return bool(sd.get("build", self.builds_default))

    def s_title(self, sd, tm):
        s = self.new_slide("ink")
        m = self.meta
        ts = m.get("title_slide") or {}
        top = " · ".join(x for x in (m.get("course"),
                                          m.get("university_short") if ts.get("university", True) else None) if x)
        if top:
            self.text(s, M, 1.5, CW, 0.5, [self.fix_text(top).upper()], 18, "signal", font=self.t["mono_font"],
                      bold=True, margin=(0, 0), name="Course")
        title = sd.get("title") or m.get("title", "Lecture")
        size, _ = fit_size([title], CW, 2.2, 56, 36, line=1.0)
        self.text(s, M, 2.15, CW, 2.3, [title], size, "on_dark", font=self.t["head_font"], bold=True, line=1.0,
                  margin=(0, 0), name="Title")
        sub = sd.get("subtitle") or m.get("subtitle")
        if sub:
            self.text(s, M, 4.55, CW, 0.6, [sub], 24, "on_dark_muted", margin=(0, 0), name="Subtitle")
        bottom = [m.get("lecturer")]
        if ts.get("university", True) and m.get("university"):
            bottom.append(m["university"])
        if ts.get("group", True):
            bottom.append(m.get("group"))
        if ts.get("year", True):
            bottom.append(m.get("academic_year") or m.get("date"))
        line1 = " · ".join(str(v) for v in bottom if v)
        extra = [v for v, on in ((m.get("website"), ts.get("website")),
                                 (ts.get("email_value"), ts.get("email"))) if v and on]
        if line1:
            self.text(s, M, 6.05 if extra else 6.3, CW, 0.45, [line1], 18, "on_dark_muted", margin=(0, 0),
                      name="Lecturer")
        if extra:
            self.text(s, M, 6.5, CW, 0.4, [" · ".join(extra)], 16, "on_dark_muted", margin=(0, 0), name="Links")
        return s, None

    def s_section(self, sd, tm):
        s = self.new_slide("ink")
        self.section_no = int(sd.get("number", self.section_no + 1))
        self.section = sd.get("title", "")
        self.text(s, M, 2.0, 3, 1.3, [f"{self.section_no:02d}"], 80, "signal", font=self.t["mono_font"], bold=True,
                  margin=(0, 0), name="Section number")
        size, _ = fit_size([self.section], CW, 1.8, 48, 34, line=1.0)
        self.text(s, M, 3.45, CW, 1.9, [self.section], size, "on_dark", font=self.t["head_font"], bold=True,
                  line=1.0, margin=(0, 0), name="Section title")
        if sd.get("subtitle"):
            self.text(s, M, 5.4, CW, 0.6, [sd["subtitle"]], 22, "on_dark_muted", margin=(0, 0))
        return s, None

    def s_break(self, sd, tm):
        s = self.new_slide("ink")
        mins = sd.get("minutes", DEFAULT_MINUTES["break"])
        self.text(s, M, 2.3, CW, 1.4, [self.s["break_title"]], 60, "on_dark", font=self.t["head_font"], bold=True,
                  margin=(0, 0))
        self.text(s, M, 3.8, CW, 0.7, [sd.get("message", self.s["break_msg"].format(f"{mins:g}"))], 26,
                  "on_dark_muted", margin=(0, 0))
        self.chip(s, M, 4.8, self.s["min_chip"].format(f"{mins:g}"))
        return s, None

    def s_objectives(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["objectives_kicker"])
        self.headline(s, sd.get("headline", self.s["objectives_head"]))
        objs = sd.get("points") or self.spec.get("objectives") or []
        if not objs:
            self.warn("no objectives (add lecture-level `objectives`)")
        if len(objs) > 4:
            self.warn(f"{len(objs)} objectives (3–4 is the target)")
        self._numbered_rows(s, objs, tm, self._build(sd))
        self.footer(s)
        return s, None

    def _numbered_rows(self, s, items, tm, build, y0=CT, y1=CB):
        n = max(len(items), 1)
        row = min(1.15, (y1 - y0) / n)
        self.check_text(*items)
        fitted = [fit_size([str(i)], CW - 1.3, row * 0.95, 26, 22, line=1.1) for i in items] or [(26, True)]
        size = min(sz for sz, _ in fitted)
        if not all(ok for _, ok in fitted):
            self.warn("numbered items are too long for their rows; shorten them")
        for i, item in enumerate(items):
            y = y0 + i * row
            num = self.text(s, M, y, 0.9, row, [f"{i + 1:02d}"], 30, "accent", font=self.t["mono_font"], bold=True,
                            margin=(0, 0), name=f"Number {i + 1}")
            txt = self.text(s, M + 1.1, y + 0.04, CW - 1.1, row, [str(item)], size, "text", line=1.1, margin=(0, 0),
                            name=f"Item {i + 1}")
            if build:
                tm.click((num, "in"), (txt, "in"))

    def s_retrieval(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["retrieval_kicker"])
        self.headline(s, sd.get("headline", self.s["retrieval_head"]))
        items = sd.get("items", [])
        n = max(len(items), 1)
        row = (CB - CT) / n
        for i, it in enumerate(items):
            y = CT + i * row
            self.text(s, M, y, 0.9, 0.6, [f"{i + 1}"], 28, "accent", font=self.t["mono_font"], bold=True,
                      margin=(0, 0))
            q, a = str(it.get("q", "")), str(it.get("a", ""))
            self.check_text(q, a)
            qs, _ = fit_size([q], CW - 1.1, row * 0.5, 24, 20)
            self.text(s, M + 1.1, y + 0.02, CW - 1.1, row * 0.5, [q], qs, "text", margin=(0, 0),
                      name=f"Question {i + 1}")
            ans = self.text(s, M + 1.1, y + row * 0.48, CW - 1.1, row * 0.45, [a], 22, "accent", margin=(0, 0),
                            name=f"Answer {i + 1}")
            tm.click((ans, "in"))
        self.footer(s)
        return s, self.s["p_retrieval"]

    def s_hook(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["hook_kicker"])
        q = str(sd.get("question") or sd.get("headline", ""))
        has_ev = bool(sd.get("evidence"))
        size, _ = fit_size([q], CW, 2.0 if has_ev else 3.5, 40, 30, line=1.1)
        self.text(s, M, 1.0, CW, 2.2 if has_ev else 4.5, [q], size, "text", font=self.t["head_font"], bold=True,
                  line=1.1, margin=(0, 0), anchor=MSO_ANCHOR.TOP if has_ev else MSO_ANCHOR.MIDDLE,
                  name="Hook question")
        if has_ev:
            self.evidence(s, sd["evidence"], M, 3.4, CW, CB - 3.4, tm, self._build(sd))
        self.footer(s, sd.get("source"))
        return s, self.s["p_hook"]

    def s_assertion(self, sd, tm):
        s = self.new_slide()
        self._content_head(s, sd)
        if not sd.get("headline"):
            self.warn("assertion slide without a headline")
        pts, ev = sd.get("points"), sd.get("evidence")
        build = self._build(sd)
        if pts and ev:
            lw = CW * 0.38
            self.bullets(s, pts, M, CT, lw, CB - CT, tm, build, size=24)
            self.evidence(s, ev, M + lw + 0.4, CT, CW - lw - 0.4, CB - CT, tm, build)
        elif ev:
            self.evidence(s, ev, M, CT, CW, CB - CT, tm, build)
        elif pts:
            self.bullets(s, pts, M, CT, CW, CB - CT, tm, build)
        else:
            self.warn("assertion slide has no evidence or points")
        self.footer(s, sd.get("source") or self.meta.get("source"))
        return s, None

    def s_code(self, sd, tm):
        s = self.new_slide()
        self._content_head(s, sd)
        steps, output = sd.get("steps"), sd.get("output")
        code = str(sd.get("code", ""))
        if steps:
            cw = CW * 0.62
            self.code_panel(s, code, M, CT, cw, CB - CT, tm, steps, sd.get("language"),
                            notes_box=(M + cw + 0.4, CT + 0.1, CW - cw - 0.4, CB - CT - 0.1))
        elif output:
            out_lines = str(output).rstrip("\n").split("\n")
            longest = max(len(l) for l in code.rstrip("\n").split("\n"))
            out_w = max(2.8, max(map(len, out_lines)) * 0.6 * 20 / 72 + 0.7)   # room for 20 pt output
            if longest <= 50 and out_w <= CW * 0.42:                  # side by side: code keeps full height
                cw = CW - out_w - 0.4
                self.code_panel(s, code, M, CT, cw, CB - CT, tm, None, sd.get("language"))
                self.ev_output(s, {"text": output}, M + cw + 0.4, CT, out_w, CB - CT, tm, True)
            else:                                                     # stacked: output under the code
                out_h = min(2.2, len(out_lines) * 0.42 + 0.5)
                used = self.code_panel(s, code, M, CT, CW, CB - CT - out_h - 0.25, tm, None, sd.get("language"))
                self.ev_output(s, {"text": output}, M, CT + used + 0.25, CW, CB - CT - used - 0.25, tm, True)
        else:
            self.code_panel(s, code, M, CT, CW, CB - CT, tm, None, sd.get("language"))
        if steps and output:
            self.warn("code slide with both steps and output: put the output on the next slide")
        self.footer(s, sd.get("source") or self.meta.get("source"))
        return s, self.s["p_code"] if steps else None

    def s_predict(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["predict_kicker"])
        self.headline(s, sd.get("question", self.s["predict_head"]))
        cw = CW * 0.58
        self.code_panel(s, sd.get("code", ""), M, CT, cw, CB - CT - 0.1, tm, None, sd.get("language"))
        rx, rw = M + cw + 0.4, CW - cw - 0.4
        y = CT
        choices = [str(c) for c in sd.get("choices") or []]
        for i, ch in enumerate(choices):
            self.text(s, rx, y, 0.55, 0.5, [chr(65 + i)], 24, "accent", font=self.t["mono_font"], bold=True,
                      margin=(0, 0))
            cs, _ = fit_size([ch], rw - 0.6, 0.55, 24, 18)
            self.text(s, rx + 0.6, y, rw - 0.6, 0.55, [ch], cs, "text", margin=(0, 0), name=f"Choice {chr(65 + i)}")
            y += 0.62
        if not choices:
            self.text(s, rx, y, rw, 0.9, [self.s["predict_blank"]], 22, "muted", italic=True, margin=(0, 0))
            y += 1.0
        ans_h = CB - y - 0.1
        card = self.rect(s, rx, y + 0.1, rw, ans_h, fill="tint", rounded=True, radius=0.05, name="Answer card")
        parts = [f"**{self.s['answer']}** {sd.get('answer', '')}"]
        if sd.get("explanation"):
            parts.append(str(sd["explanation"]))
        self.check_text(*parts)
        asz, fits = fit_size(parts, rw - 0.5, ans_h - 0.4, 22, 18, gap_pt=8)
        if not fits:
            self.warn("predict answer/explanation too long for the card")
        self.text(card, 0, 0, 0, 0, parts, asz, "text", gap=8, margin=(0.25, 0.2), anchor=MSO_ANCHOR.TOP)
        tm.click((card, "in"))
        self.footer(s, sd.get("source") or self.meta.get("source"))
        proto = self.s["p_predict"]
        if sd.get("misconception"):
            proto += "\n" + self.s["p_misconception"].format(sd["misconception"])
        return s, proto

    def s_quiz(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["quiz_kicker"])
        self.headline(s, sd.get("question", ""))
        choices = [str(c) for c in sd.get("choices", [])]
        self.check_text(*choices)
        correct = int(sd.get("correct", 0))
        if len(choices) < 3:
            self.warn("quiz has fewer than 3 choices")
        cols = 2 if len(choices) > 2 else 1
        rows = math.ceil(len(choices) / cols) if choices else 1
        gap = 0.3
        expl_h = 1.0 if sd.get("explanation") else 0
        cw = (CW - gap * (cols - 1)) / cols
        ch = min(1.3, (CB - CT - expl_h - gap * rows) / rows)
        sizes = [fit_size([c], cw - 1.2, ch - 0.2, 24, 18)[0] for c in choices] or [24]
        size = min(sizes)
        hit = None
        for i, c in enumerate(choices):
            r, k = divmod(i, cols)
            x, y = M + k * (cw + gap), CT + r * (ch + gap)
            self.rect(s, x, y, cw, ch, fill="tint", rounded=True, radius=0.08, name=f"Choice {chr(65 + i)}")
            self.text(s, x + 0.25, y, 0.7, ch, [chr(65 + i)], 30, "accent", font=self.t["mono_font"], bold=True,
                      anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0))
            self.text(s, x + 0.95, y, cw - 1.15, ch, [c], size, "text", anchor=MSO_ANCHOR.MIDDLE, margin=(0, 0))
            if i == correct:
                hit = self.rect(s, x - 0.04, y - 0.04, cw + 0.08, ch + 0.08, fill="accent", alpha=18, line="accent",
                                line_w=4, rounded=True, radius=0.08, name="Correct answer")
        expl = None
        if sd.get("explanation"):
            ey = CT + rows * (ch + gap)
            es, _ = fit_size([sd["explanation"]], CW, CB - ey, 22, 18)
            self.check_text(sd["explanation"])
            expl = self.text(s, M, ey, CW, CB - ey, [f"**{self.s['answer']}** {sd['explanation']}"], es, "text",
                             margin=(0, 0), name="Explanation")
        tm.click((hit, "in") if hit else None, (expl, "in") if expl else None)
        self.footer(s, sd.get("source") or self.meta.get("source"))
        proto = self.s["p_quiz"]
        if sd.get("distractors"):
            proto += "\n" + self.s["p_distractors"].format(sd["distractors"])
        return s, proto

    def _mode_slide(self, sd, label):
        s = self.new_slide("mode")
        self.text(s, M, KICKER_Y, CW, 0.36, [self.fix_text(label).upper()], 14, "signal", font=self.t["mono_font"],
                  bold=True, margin=(0, 0), name="Kicker")
        title = str(sd.get("title", ""))
        size, _ = fit_size([title], CW - 2.2, HEAD_H, 36, 28, line=1.05)
        self.text(s, M, HEAD_Y, CW - 2.2, HEAD_H, [title], size, "on_dark", font=self.t["head_font"], bold=True,
                  line=1.05, margin=(0, 0), name="Headline")
        mins = sd.get("minutes", DEFAULT_MINUTES.get(sd.get("type"), 10))
        self.chip(s, W - M - 1.7, HEAD_Y + 0.05, self.s["min_chip"].format(f"{mins:g}"))
        return s

    def s_livecode(self, sd, tm):
        s = self._mode_slide(sd, self.s["livecode_kicker"])
        goals = [str(g) for g in sd.get("goals") or []]
        self.check_text(*goals)
        if sd.get("file"):
            self.text(s, M, CT - 0.1, CW, 0.5, [self.s["starter"].format(sd["file"])], 18, "on_dark_muted",
                      font=self.t["mono_font"], margin=(0, 0), name="Starter file")
        gsize, _ = fit_size(goals, CW - 0.6, CB - CT - 0.6, 26, 22, gap_pt=12)
        shp = self.text(s, M, CT + 0.55, CW, CB - CT - 0.6, goals, gsize, "on_dark", bullets=True,
                        gap=12, name="Goals", code_color="signal", marker="signal")
        for i in range(len(goals)):
            tm.click((shp, "in", i))
        self.footer(s, sd.get("source"), dark=True)
        files = {f.get("name") for f in (self.spec.get("livecode") or {}).get("files", [])}
        if sd.get("file") and sd["file"] not in files:
            self.warn(f"livecode file '{sd['file']}' is not defined in lecture `livecode.files`")
        return s, self.s["p_livecode"]

    def s_activity(self, sd, tm):
        label = self.s["activity_kicker"] + (self.s["activity_pairs"] if sd.get("pairs") else "")
        s = self._mode_slide(sd, label)
        steps = [str(x) for x in sd.get("steps", [])]
        self.check_text(*steps)
        top = CT
        if sd.get("file"):
            self.text(s, M, CT - 0.1, CW, 0.5, [self.s["open_file"].format(sd["file"])], 18, "on_dark_muted",
                      font=self.t["mono_font"], margin=(0, 0), name="Starter file")
            top = CT + 0.55
        size, fits = fit_size(steps, CW - 0.6, CB - top, 26, 22, gap_pt=12)
        if not fits:
            self.warn("activity steps too long")
        self.text(s, M, top, CW, CB - top, steps, size, "on_dark", numbered=True, gap=12, name="Steps",
                  code_color="signal", marker="signal")
        self.footer(s, None, dark=True)
        proto = self.s["p_activity"]
        if sd.get("hint"):
            proto += "\n" + self.s["p_hint"].format(sd["hint"])
        return s, proto

    def s_compare(self, sd, tm):
        s = self.new_slide()
        self._content_head(s, sd)
        gap = 0.5
        cw = (CW - gap) / 2
        build = self._build(sd)
        for k, side in enumerate(("left", "right")):
            col = sd.get(side) or {}
            x = M + k * (cw + gap)
            card = self.rect(s, x, CT, cw, CB - CT, fill="tint", rounded=True, radius=0.04, name=f"{side} column")
            title = self.text(s, x + 0.3, CT + 0.2, cw - 0.6, 0.6, [col.get("title", "")], 26,
                              "accent", font=self.t["head_font"], bold=True, margin=(0, 0), name=f"{side} title")
            if build:
                tm.click((card, "in"), (title, "in"))
            body_y = CT + 0.95
            if col.get("evidence"):
                self.evidence(s, col["evidence"], x + 0.3, body_y, cw - 0.6, CB - body_y - 0.2, tm, build)
            else:
                self.bullets(s, col.get("points", []), x + 0.3, body_y, cw - 0.6, CB - body_y - 0.2, tm, build,
                             size=24)
        self.footer(s, sd.get("source") or self.meta.get("source"))
        return s, None

    def s_recap(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["recap_kicker"])
        self.headline(s, sd.get("headline", self.s["recap_head"]))
        pts = sd.get("points") or []
        if not 3 <= len(pts) <= 5:
            self.warn(f"recap has {len(pts)} points (3–5)")
        self._numbered_rows(s, pts, tm, self._build(sd))
        self.footer(s)
        return s, self.s["p_recap"]

    def s_exit(self, sd, tm):
        s = self.new_slide()
        self.kicker(s, sd.get("kicker") or self.s["exit_kicker"])
        qs = [str(q) for q in sd.get("questions", [])]
        self.check_text(*qs)
        self.headline(s, sd.get("headline", self.s["exit_head"]))
        size, _ = fit_size(qs, CW - 0.6, CB - CT - 1.0, 28, 22, gap_pt=14)
        self.text(s, M, CT, CW, CB - CT - 1.0, qs, size, "text", numbered=True, gap=14, name="Exit questions")
        if sd.get("next"):
            self.text(s, M, CB - 0.6, CW, 0.5, [sd["next"]], 20, "muted", italic=True, margin=(0, 0), name="Next")
        self.footer(s)
        return s, self.s["p_exit"]

    # ---- run -----------------------------------------------------------------
    def build(self):
        for sd in self.spec.get("slides", []):
            self.add(sd)
        dur = float(self.meta.get("duration_min", 80))
        if abs(self.clock - dur) > 5:
            self.warnings.append(f"lecture: slides add up to {self.clock:g} min but the lecture is {dur:g} min "
                                 f"(adjust `minutes` or content)")
        for o in self.spec.get("objectives") or []:
            if re.match(r"\W*(tushun|bil(ib|ish)|o'rgan|understand|know|learn)", plain(o), re.I):
                self.warnings.append("lecture: objective not measurable (use yoza oladi / tushuntira oladi / "
                                     f"kuzata oladi / taqqoslay oladi…): “{plain(o)[:50]}”")
        types = [t for _, t, *_ in self.timeline]
        for must in ("objectives", "recap"):
            if must not in types:
                self.warnings.append(f"lecture: no `{must}` slide")
        if not ACTIVE_TYPES & set(types):
            self.warnings.append("lecture: no active-learning slides at all")
        return self.prs


# --------------------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec", type=Path)
    ap.add_argument("--deck", type=Path, help="output .pptx (overrides meta.deck_file)")
    ap.add_argument("--only", default="deck,code", help="comma list of deck,code")
    ap.add_argument("--lint-only", action="store_true", help="print timeline + warnings, write nothing")
    ap.add_argument("--force", action="store_true", help="overwrite an existing deck")
    args = ap.parse_args()

    spec_path = args.spec.resolve()
    spec = load_spec(spec_path)
    meta = spec.setdefault("meta", {})
    slug = meta.get("slug") or slugify(meta.get("title", spec_path.stem))
    base = spec_path.parent
    deck_path = (args.deck.resolve() if args.deck else
                 (base / meta["deck_file"]).resolve() if meta.get("deck_file") else base / "build" / f"{slug}.pptx")
    want = {w.strip() for w in args.only.split(",")}

    deck = LectureDeck(spec)
    prs = deck.build()

    cfg = spec.get("_config")
    print(f"Config: {cfg if cfg else 'none found (using defaults)'}")
    print(f"Timeline ({deck.clock:g} of {meta.get('duration_min', 80)} min):")
    for n, typ, a, b, clicks, label in deck.timeline:
        print(f"  {n:>3}  {a:>5g}–{b:<5g} {typ:<11} {clicks:>2} clicks  {label[:60]}")
    if deck.warnings:
        print(f"\n{len(deck.warnings)} warning(s):")
        for w in deck.warnings:
            print("  - " + w)
    else:
        print("\nNo design-rule warnings.")
    if args.lint_only:
        return

    written = []
    if "deck" in want:
        if deck_path.exists() and not args.force:
            sys.exit(f"\nNOT WRITTEN: {deck_path} already exists. It may contain hand edits.\n"
                     "Ask the teacher, then rerun with --force, or build to another file with --deck.")
        deck_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(deck_path)
        written.append(deck_path)
    if "code" in want and spec.get("livecode"):
        written += companions.livecode_files(spec, base / "livecode")
    (base / "build").mkdir(exist_ok=True)
    (base / "build" / "lint.txt").write_text("\n".join(deck.warnings) + "\n", encoding="utf-8")
    print("\nWrote:")
    for p in written:
        print(f"  {p}")


if __name__ == "__main__":
    main()
